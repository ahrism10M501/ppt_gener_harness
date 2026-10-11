"""Font resolution, fixed text boxes and LibreOffice rendering shared by the build."""

import math
import os
from pathlib import Path
import shutil
import subprocess
import tempfile

import pymupdf as fitz
from lxml import etree
from PIL import Image, ImageDraw
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt

from .template import NS


def resolve_fonts(tokens):
    profile = {"verification": "fontconfig" if shutil.which("fc-match") else "unverified"}
    for script in ("latin", "korean"):
        preferred = tokens["fonts"][script]
        fallback = tokens["fonts"][f"fallback_{script}"]
        resolved = preferred
        if profile["verification"] == "fontconfig":
            for candidate in (preferred, fallback):
                family = subprocess.run(["fc-match", "-f", "%{family}", candidate], check=True, capture_output=True, text=True).stdout.strip()
                if candidate.casefold() in family.casefold().split(","):
                    resolved = candidate
                    break
            else:
                raise RuntimeError(f"Missing required {script} font: {preferred} / {fallback}")
        profile[script] = {"requested": preferred, "resolved": resolved, "fallback": resolved != preferred}
    return profile


def text(shape, value, size, color, profile, *, center=False, bold=False, bullets=False):
    frame = shape.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = True
    # CJK font ascenders extend above LibreOffice's nominal first-line origin.
    frame.margin_top = Pt(profile["text_top_inset_pt"])
    frame.margin_bottom = 0
    frame.margin_left = frame.margin_right = 0
    for i, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
        p.line_spacing = 1.0
        p.space_before = Pt(0)
        p.space_after = Pt(14 if bullets else 0)
        ppr = p._p.get_or_add_pPr()
        for child in list(ppr):
            if etree.QName(child).localname in ("buNone", "buChar", "buAutoNum"):
                ppr.remove(child)
        if bullets:
            ppr.set("marL", "228600")
            ppr.set("indent", "-152400")
            etree.SubElement(ppr, f"{{{NS['a']}}}buChar", char="▪")
        else:
            etree.SubElement(ppr, f"{{{NS['a']}}}buNone")
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
        r.font.name = profile["latin"]["resolved"]
        props = r._r.get_or_add_rPr()
        ea = props.find(f"{{{NS['a']}}}ea")
        if ea is None:
            ea = etree.SubElement(props, f"{{{NS['a']}}}ea")
        ea.set("typeface", profile["korean"]["resolved"])


def fit_picture(image_size, box):
    x, y, w, h = box
    scale = min(w / image_size[0], h / image_size[1])
    width, height = image_size[0] * scale, image_size[1] * scale
    return x + (w - width) / 2, y + (h - height) / 2, width, height


def libreoffice_binary():
    for name in ("libreoffice", "soffice"):
        if path := shutil.which(name):
            return Path(path)
    candidates = [Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")]
    for variable in ("PROGRAMFILES", "PROGRAMFILES(X86)"):
        if root := os.environ.get(variable):
            candidates.append(Path(root) / "LibreOffice/program/soffice.exe")
    for path in candidates:
        if path.is_file():
            return path
    raise RuntimeError("LibreOffice is required; see docs/setup.md")


def chrome_binary():
    """Mermaid browser: PUPPETEER_EXECUTABLE_PATH, PATH, then the macOS and Windows install folders; None if absent."""
    if path := os.environ.get("PUPPETEER_EXECUTABLE_PATH"):
        return path
    for name in ("google-chrome", "chromium", "chromium-browser"):
        if path := shutil.which(name):
            return path
    candidates = [Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")]
    for variable in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA"):
        if root := os.environ.get(variable):
            candidates.append(Path(root) / "Google/Chrome/Application/chrome.exe")
    return next((str(path) for path in candidates if path.is_file()), None)


def render_pdf(source, destination):
    binary = libreoffice_binary()
    staging_parent = None
    if str(binary).startswith("/snap/"):
        staging_parent = Path.home() / "snap/libreoffice/common"
        staging_parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="rvppt-preview-", dir=staging_parent) as directory:
        stage = Path(directory)
        local = stage / "input.pptx"
        shutil.copy2(source, local)
        profile = (stage / "lo-profile").as_uri()
        result = subprocess.run([str(binary), f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf", "--outdir", str(stage), str(local)], capture_output=True, text=True, timeout=120)
        generated = stage / "input.pdf"
        if result.returncode or not generated.is_file() or generated.stat().st_size == 0:
            raise RuntimeError(f"LibreOffice conversion failed: {result.stdout} {result.stderr}")
        with fitz.open(generated) as doc:
            if len(doc) == 0:
                raise RuntimeError("LibreOffice generated an empty PDF")
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(generated, destination)


def contact_sheet(images, labels, output):
    columns, width, title_h = 2, 640, 28
    tile_h = round(width * 3 / 4) + title_h
    sheet = Image.new("RGB", (width * columns, tile_h * math.ceil(len(images) / columns)), "#e5e7eb")
    draw = ImageDraw.Draw(sheet)
    for idx, (path, label) in enumerate(zip(images, labels)):
        x, y = idx % columns * width, idx // columns * tile_h
        draw.text((x + 12, y + 8), label, fill="black")
        with Image.open(path) as picture:
            sheet.paste(picture.resize((width, tile_h - title_h)), (x, y + title_h))
    sheet.save(output)
