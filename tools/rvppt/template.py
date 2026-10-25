"""Extract the original template and append fixed RVLAB/Pandoc or multi-preset layouts."""

import argparse
from copy import deepcopy
import hashlib
import json
import posixpath
from pathlib import Path
import re
import shutil
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from lxml import etree
from pptx import Presentation
import yaml

ROOT = Path(__file__).resolve().parents[2]
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
EMU = 914400
PRESET_ID_PATTERN = re.compile(r"[a-z0-9][a-z0-9-]{0,40}\Z")


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def design_path(preset="rvlab", root=ROOT):
    """Each template is one folder, templates/<id>/, whose DESIGN.md holds its tokens and style conventions."""
    if not PRESET_ID_PATTERN.fullmatch(str(preset)):
        raise ValueError(f"Invalid preset ID: {preset}")
    return Path(root) / "templates" / preset / "DESIGN.md"


def read_design(path):
    """The first ```yaml block of a DESIGN.md is the token set; the prose below it is for the agent."""
    text = Path(path).read_text(encoding="utf-8")
    return yaml.safe_load(text.split("```yaml\n", 1)[1].split("```", 1)[0] if "```yaml\n" in text else text)


def task_preset(path):
    """deck.preset of the task folder that contains path (a figure or script inside tasks/<task>/), else rvlab."""
    for folder in [Path(path).resolve(), *Path(path).resolve().parents]:
        slides = folder / "slides.md"
        if slides.is_file():
            text = slides.read_text(encoding="utf-8")
            meta = yaml.safe_load(text.split("---", 2)[1]) if text.startswith("---") else None
            deck = meta.get("deck") if isinstance(meta, dict) else None
            preset = deck.get("preset") if isinstance(deck, dict) else None
            return preset.strip() if isinstance(preset, str) and preset.strip() else "rvlab"
    return "rvlab"


def load_tokens(design=None, preset=None, root=ROOT):
    path = Path(design) if design else design_path(preset or "rvlab", root)
    if not path.is_file():
        raise ValueError(f"Unknown design preset: {preset} (run `pptgen presets` to list available presets)")
    tokens = read_design(path)
    if not isinstance(tokens, dict) or tokens.get("schema_version") != 1:
        raise ValueError("Unsupported design token schema")
    tokens.setdefault("preset_id", preset or tokens.get("preset_id", "rvlab"))
    width_in = float(tokens["slide"]["width_in"])
    height_in = float(tokens["slide"]["height_in"])
    for key in ("title", "header_rule", "body"):
        x, y, w, h = tokens["geometry_in"][key]
        if min(x, y) < 0 or min(w, h) <= 0 or x + w > width_in + 1e-3 or y + h > height_in + 1e-3:
            raise ValueError(f"Invalid fixed geometry: {key}")
    colors = {**tokens["palette"], **{k: v for k, v in tokens.get("table", {}).items() if k not in ("rule_pt", "first_column_bold")}}
    # YAML reads an unquoted all-digit hex such as 111833 as a number, which would silently drop the color.
    unquoted = [k for k, v in colors.items() if not isinstance(v, str)]
    if unquoted:
        raise ValueError(f"Quote these color values in {path.name}: {', '.join(unquoted)}")
    return tokens


def semantic_palette(preset="rvlab"):
    """Return shared visual roles as #RRGGBB, derived from the preset tokens."""
    return palette_roles(load_tokens(preset=preset))


def palette_roles(tokens):
    palette = tokens["palette"]
    roles = {
        role: "#" + palette[key].lstrip("#").upper()
        for role, key in {
            "accent": "primary", "band": "band", "secondary": "secondary",
            "ink": "text", "section_accent": "section_accent", "background": "background",
        }.items()
    }

    def blend(color, factor):
        background = roles["background"]
        channels = (
            round(int(color[i:i + 2], 16) * (1 - factor) + int(background[i:i + 2], 16) * factor)
            for i in (1, 3, 5)
        )
        return "#" + "".join(f"{channel:02X}" for channel in channels)

    # These blends retain the original RVLAB chart tint and neutral exactly.
    roles["tint"] = blend(roles["secondary"], 0.8)
    roles["neutral"] = blend(roles["ink"], 0.65)
    return roles


def list_presets(root=ROOT):
    presets_dir = Path(root) / "templates"
    items = []
    if presets_dir.is_dir():
        for path in sorted(presets_dir.glob("*/DESIGN.md")):
            data = read_design(path)
            if not isinstance(data, dict):
                continue
            ref = Path(root) / data.get("reference", "")
            items.append({
                "preset_id": data.get("preset_id", path.parent.name),
                "name": data.get("name", path.parent.name),
                "description": data.get("description", ""),
                "mode": data.get("mode", "preset"),
                "aspect": data["slide"]["aspect"],
                "slide_in": [data["slide"]["width_in"], data["slide"]["height_in"]],
                "palette": data["palette"],
                "reference": data.get("reference"),
                "reference_ready": ref.is_file(),
                # Tokens plus the character and style conventions an agent loads before working on a deck.
                "design": str(path.relative_to(root)),
            })
    return {"presets": items, "count": len(items)}


def slots(tokens, count, body=None):
    if count not in (1, 2, 3, 4):
        raise ValueError("RVLAB supports 1/2/3/4 figure slots")
    x, y, w, h = body or tokens["geometry_in"]["body"]
    gap = tokens["geometry_in"]["gap"]
    columns, rows = (2, 2) if count == 4 else (count, 1)
    sw, sh = (w - gap * (columns - 1)) / columns, (h - gap * (rows - 1)) / rows
    return [(x + c * (sw + gap), y + r * (sh + gap), sw, sh)
            for r in range(rows) for c in range(columns)]


def parse(data):
    return etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))


def xml(root):
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def extract_manifest(source):
    prs = Presentation(source)
    result = {
        "schema_version": 1,
        "source": {"path": str(Path(source).resolve().relative_to(ROOT)) if Path(source).resolve().is_relative_to(ROOT) else Path(source).name, "sha256": sha256(source)},
        "slide_emu": [prs.slide_width, prs.slide_height],
        "slide_inches": [prs.slide_width / EMU, prs.slide_height / EMU],
        "masters": len(prs.slide_masters), "specimen_slides": len(prs.slides),
        "layouts": [],
    }
    for layout in prs.slide_layouts:
        shapes = []
        for shape in layout.shapes:
            record = {"id": shape.shape_id, "name": shape.name,
                      "box_emu": [shape.left, shape.top, shape.width, shape.height]}
            if shape.is_placeholder:
                record["placeholder"] = {"idx": shape.placeholder_format.idx,
                                          "type": str(shape.placeholder_format.type)}
            if shape.has_text_frame:
                record["text"] = shape.text
            shapes.append(record)
        result["layouts"].append({"name": layout.name, "shapes": shapes})
    with ZipFile(source) as z:
        layouts = [parse(z.read(n)) for n in z.namelist() if n.startswith("ppt/slideLayouts/") and n.endswith(".xml")]
        result["layout_fonts"] = sorted({v for root in layouts for v in root.xpath(".//a:latin/@typeface | .//a:ea/@typeface", namespaces=NS)})
        result["layout_direct_colors"] = sorted({v for root in layouts for v in root.xpath(".//a:srgbClr/@val", namespaces=NS)})
        result["themes"] = {
            n: {etree.QName(e).localname: [dict(c.attrib) for c in e]
                for e in parse(z.read(n)).xpath(".//a:clrScheme/*", namespaces=NS)}
            for n in z.namelist() if n.startswith("ppt/theme/") and n.endswith(".xml")
        }
    return result


def set_box(shape, box):
    xfrm = shape.find("p:spPr/a:xfrm", NS)
    if xfrm is None:  # Placeholders that inherit their position get an explicit one.
        xfrm = etree.Element(f"{{{NS['a']}}}xfrm")
        shape.find("p:spPr", NS).insert(0, xfrm)
        etree.SubElement(xfrm, f"{{{NS['a']}}}off"); etree.SubElement(xfrm, f"{{{NS['a']}}}ext")
    off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
    values = [str(round(v * EMU)) for v in box]
    off.set("x", values[0]); off.set("y", values[1])
    ext.set("cx", values[2]); ext.set("cy", values[3])


def placeholder(shape):
    return shape.find("p:nvSpPr/p:nvPr/p:ph", NS)


def set_text_style(shape, size, color, tokens, *, center=False, bullets=False):
    body = shape.find("p:txBody", NS)
    body_pr = body.find("a:bodyPr", NS)
    body_pr.attrib.update({"lIns": "0", "rIns": "0", "tIns": str(round(tokens["text_top_inset_pt"] * 12700)), "bIns": "0", "anchor": "t", "wrap": "square"})
    for child in list(body_pr):
        body_pr.remove(child)
    etree.SubElement(body_pr, f"{{{NS['a']}}}noAutofit")
    style = body.find("a:lstStyle", NS)
    style.clear()
    for level in range(1, 10):
        p = etree.SubElement(style, f"{{{NS['a']}}}lvl{level}pPr", algn="ctr" if center else "l")
        if bullets:
            p.set("marL", str(228600 * min(level, 2)))
            p.set("indent", "-152400")
        spacing = etree.SubElement(p, f"{{{NS['a']}}}lnSpc")
        etree.SubElement(spacing, f"{{{NS['a']}}}spcPct", val="100000")
        if bullets:
            after = etree.SubElement(p, f"{{{NS['a']}}}spcAft")
            etree.SubElement(after, f"{{{NS['a']}}}spcPts", val="1000")
            etree.SubElement(p, f"{{{NS['a']}}}buChar", char="▪" if level == 1 else "•")
        else:
            etree.SubElement(p, f"{{{NS['a']}}}buNone")
        r = etree.SubElement(p, f"{{{NS['a']}}}defRPr", sz=str(int((size if level == 1 else min(size, tokens['sizes_pt']['secondary'])) * 100)))
        fill = etree.SubElement(r, f"{{{NS['a']}}}solidFill")
        etree.SubElement(fill, f"{{{NS['a']}}}srgbClr", val=color)
        etree.SubElement(r, f"{{{NS['a']}}}latin", typeface=tokens["fonts"]["latin"])
        etree.SubElement(r, f"{{{NS['a']}}}ea", typeface=tokens["fonts"]["korean"])
    for paragraph in body.findall("a:p", NS):
        body.remove(paragraph)
    etree.SubElement(body, f"{{{NS['a']}}}p")
    shape.find("p:nvSpPr/p:cNvSpPr", NS).attrib.pop("txBox", None)


def content_layout(base, name, count, tokens):
    root = deepcopy(base)
    root.set("type", "obj"); root.set("preserve", "1")
    root.find("p:cSld", NS).set("name", name)
    tree = root.find("p:cSld/p:spTree", NS)
    shapes = tree.findall("p:sp", NS)
    title = next(s for s in shapes if placeholder(s) is not None and placeholder(s).get("type") == "title")
    body = next(s for s in shapes if placeholder(s) is not None and placeholder(s).get("idx") == "1")
    # The white title box masks the rule behind it; RVLAB decks widen it to this fixed tab.
    set_box(title, tokens["geometry_in"]["title"])
    tree.remove(body)
    for idx, box in enumerate(slots(tokens, count), 1):
        cell = deepcopy(body)
        ph = placeholder(cell)
        ph.attrib.clear(); ph.set("idx", str(idx))
        # Pandoc recognizes object placeholders only when the type attribute is absent.
        props = cell.find("p:nvSpPr/p:cNvPr", NS)
        props.set("id", str(100 + idx)); props.set("name", f"RVLAB Content {idx}")
        set_box(cell, box)
        set_text_style(cell, tokens["sizes_pt"]["headline" if count == 1 else "columns"], tokens["palette"]["text"], tokens, bullets=True)
        tree.append(cell)
    if name == "Blank":
        for s in list(tree.findall("p:sp", NS)):
            if placeholder(s) is not None and placeholder(s).get("type") not in ("sldNum", "ftr"):
                tree.remove(s)
    return root


def validate_package(path):
    """Check internal relationships and XML before replacing a good reference."""
    with ZipFile(path) as z:
        names = set(z.namelist())
        if z.testzip() is not None:
            raise ValueError("Corrupt PPTX archive")
        for name in names:
            if name.endswith((".xml", ".rels")):
                root = parse(z.read(name))
                if name.endswith(".rels"):
                    base = posixpath.dirname(posixpath.dirname(name))
                    for rel in root:
                        if rel.get("TargetMode") == "External":
                            continue
                        target = rel.get("Target", "")
                        target = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                        if target not in names:
                            raise ValueError(f"Missing PPTX relationship target: {name} → {target}")
        for override in parse(z.read("[Content_Types].xml")):
            if "PartName" in override.attrib and override.get("PartName").lstrip("/") not in names:
                raise ValueError("Missing content type part")
    return True


def _transform_xml_for_preset(root_el, tokens, sx, sy, *, strip_pics=True):
    palette = tokens["palette"]
    color_map = {
        "0070C0": palette["primary"].upper(),
        "053F95": palette["band"].upper(),
        "002060": palette["band"].upper(),
        "5B9BD5": palette["secondary"].upper(),
        "FFFF00": palette["section_accent"].upper(),
        "000000": palette["text"].upper(),
    }
    if strip_pics:
        for pic in list(root_el.xpath(".//p:pic", namespaces=NS)):
            parent = pic.getparent()
            if parent is not None:
                parent.remove(pic)
    if abs(sx - 1.0) > 1e-6 or abs(sy - 1.0) > 1e-6:
        for xfrm in root_el.xpath(".//p:sp/p:spPr/a:xfrm", namespaces=NS):
            off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
            if off is not None and "x" in off.attrib and "y" in off.attrib:
                off.set("x", str(round(int(off.get("x")) * sx)))
                off.set("y", str(round(int(off.get("y")) * sy)))
            if ext is not None and "cx" in ext.attrib and "cy" in ext.attrib:
                ext.set("cx", str(round(int(ext.get("cx")) * sx)))
                ext.set("cy", str(round(int(ext.get("cy")) * sy)))
    for clr in root_el.xpath(".//a:srgbClr", namespaces=NS):
        val = clr.get("val", "").upper()
        if val in color_map:
            clr.set("val", color_map[val])


RVLAB_ORIGINAL = "templates/rvlab/RVLAB_PPT양식.pptx"

# Placeholders a theme may restyle, keyed by (type, idx) in the RVLAB-derived layouts.
STYLE_PLACEHOLDERS = {
    "cover": {("ctrTitle", None): "title", ("subTitle", "1"): "subtitle", ("body", "2"): "presenter", ("body", "3"): "affiliation", ("body", "4"): "date"},
    "toc": {("body", "1"): "body"},
    "section": {("ctrTitle", None): "title", ("subTitle", "1"): "subtitle"},
    "content": {("title", None): "title", ("sldNum", "12"): "number"},
}
SHAPES = {"rect": "rect", "round": "roundRect", "ellipse": "ellipse", "line": "line", "text": "rect"}


def color_of(value, tokens):
    """Palette key (primary, band, ...), semantic role (accent, tint, neutral, ink, ...) or literal RRGGBB -> RRGGBB."""
    if value in tokens["palette"]:
        return tokens["palette"][value].lstrip("#").upper()
    roles = palette_roles(tokens)
    return roles[value].lstrip("#").upper() if value in roles else str(value).lstrip("#").upper()


_color = color_of


def _solid(parent, color):
    fill = etree.SubElement(parent, f"{{{NS['a']}}}solidFill")
    etree.SubElement(fill, f"{{{NS['a']}}}srgbClr", val=color)
    return fill


def _run_style(parent, spec, tokens, tag="defRPr"):
    fonts = tokens["fonts"]
    heading = spec.get("font", "body") == "heading"
    run = etree.SubElement(parent, f"{{{NS['a']}}}{tag}", sz=str(round(spec.get("size", 18) * 100)), b="1" if spec.get("bold") else "0")
    if spec.get("spacing"):
        run.set("spc", str(round(spec["spacing"] * 100)))
    _solid(run, _color(spec.get("color", "text"), tokens))
    etree.SubElement(run, f"{{{NS['a']}}}latin", typeface=fonts.get("heading_latin", fonts["latin"]) if heading else fonts["latin"])
    etree.SubElement(run, f"{{{NS['a']}}}ea", typeface=fonts.get("heading_korean", fonts["korean"]) if heading else fonts["korean"])


def _style_text(body, spec, tokens):
    """Rewrite a txBody's frame and level styles from a theme spec (levels 1-2 for contents lists)."""
    body_pr = body.find("a:bodyPr", NS)
    body_pr.attrib.update({"lIns": "0", "rIns": "0", "tIns": "0", "bIns": "0", "wrap": "square", "anchor": spec.get("anchor", "t")})
    for child in list(body_pr):
        body_pr.remove(child)
    etree.SubElement(body_pr, f"{{{NS['a']}}}noAutofit")
    style = body.find("a:lstStyle", NS)
    if style is None:
        style = etree.Element(f"{{{NS['a']}}}lstStyle"); body_pr.addnext(style)
    style.clear()
    for level in (1, 2):
        current = dict(spec, **spec.get("level2", {})) if level == 2 else spec
        p = etree.SubElement(style, f"{{{NS['a']}}}lvl{level}pPr", algn=current.get("align", "l"))
        if spec.get("bullets"):
            p.set("marL", str(228600 * level)); p.set("indent", "-228600")
        spacing = etree.SubElement(p, f"{{{NS['a']}}}lnSpc")
        etree.SubElement(spacing, f"{{{NS['a']}}}spcPct", val=str(round(current.get("line", 1.0) * 100000)))
        after = etree.SubElement(p, f"{{{NS['a']}}}spcAft")
        etree.SubElement(after, f"{{{NS['a']}}}spcPts", val=str(round(current.get("after", 0) * 100)))
        if spec.get("bullets"):
            etree.SubElement(p, f"{{{NS['a']}}}buChar", char=spec["bullets"][level - 1])
        else:
            etree.SubElement(p, f"{{{NS['a']}}}buNone")
        _run_style(p, current, tokens)


def _decor(spec, index, tokens):
    """A user-drawn layout shape: rect, round, ellipse, line or fixed text, behind the placeholders."""
    p, a = NS["p"], NS["a"]
    sp = etree.Element(f"{{{p}}}sp")
    nv = etree.SubElement(sp, f"{{{p}}}nvSpPr")
    etree.SubElement(nv, f"{{{p}}}cNvPr", id=str(900 + index), name=f"Theme decor {index}")
    etree.SubElement(nv, f"{{{p}}}cNvSpPr", txBox="1") if spec["shape"] == "text" else etree.SubElement(nv, f"{{{p}}}cNvSpPr")
    etree.SubElement(nv, f"{{{p}}}nvPr", userDrawn="1")
    pr = etree.SubElement(sp, f"{{{p}}}spPr")
    xfrm = etree.SubElement(pr, f"{{{a}}}xfrm")
    x, y, w, h = spec["box"]
    etree.SubElement(xfrm, f"{{{a}}}off", x=str(round(x * EMU)), y=str(round(y * EMU)))
    etree.SubElement(xfrm, f"{{{a}}}ext", cx=str(round(w * EMU)), cy=str(round(h * EMU)))
    geom = etree.SubElement(pr, f"{{{a}}}prstGeom", prst=SHAPES[spec["shape"]])
    etree.SubElement(geom, f"{{{a}}}avLst")
    if spec.get("fill") and spec["shape"] != "line":
        _solid(pr, _color(spec["fill"], tokens))
    else:
        etree.SubElement(pr, f"{{{a}}}noFill")
    line = etree.SubElement(pr, f"{{{a}}}ln", w=str(round(spec.get("line_pt", 1) * 12700)))
    if spec.get("line"):
        _solid(line, _color(spec["line"], tokens))
    else:
        etree.SubElement(line, f"{{{a}}}noFill")
    if spec["shape"] == "text":
        body = etree.SubElement(sp, f"{{{p}}}txBody")
        etree.SubElement(body, f"{{{a}}}bodyPr")
        etree.SubElement(body, f"{{{a}}}lstStyle")
        _style_text(body, spec, tokens)
        paragraph = etree.SubElement(body, f"{{{a}}}p")
        run = etree.SubElement(paragraph, f"{{{a}}}r")
        _run_style(run, spec, tokens, tag="rPr")
        etree.SubElement(run, f"{{{a}}}t").text = spec["text"]
    return sp


def set_background(root, color):
    c_sld = root.find("p:cSld", NS)
    for old in c_sld.findall("p:bg", NS):
        c_sld.remove(old)
    bg = etree.Element(f"{{{NS['p']}}}bg")
    pr = etree.SubElement(bg, f"{{{NS['p']}}}bgPr")
    _solid(pr, color)
    etree.SubElement(pr, f"{{{NS['a']}}}effectLst")
    c_sld.insert(0, bg)


def apply_layout_style(root, role, tokens):
    """Replace the RVLAB decoration of one layout with the theme's: placeholders restyled, decor redrawn."""
    spec = tokens["style"].get(role, {})
    tree = root.find("p:cSld/p:spTree", NS)
    for shape in list(tree):
        if etree.QName(shape).localname in ("sp", "pic", "grpSp", "cxnSp") and placeholder(shape) is None:
            tree.remove(shape)
    for shape in tree.findall("p:sp", NS):
        ph = placeholder(shape)
        key = STYLE_PLACEHOLDERS[role].get((ph.get("type"), ph.get("idx")))
        if key is None or key not in spec:
            continue
        pr = shape.find("p:spPr", NS)
        for child in list(pr):
            if etree.QName(child).localname in ("solidFill", "noFill", "gradFill", "ln"):
                pr.remove(child)
        etree.SubElement(pr, f"{{{NS['a']}}}noFill")
        if "box" in spec[key]:
            set_box(shape, spec[key]["box"])
        _style_text(shape.find("p:txBody", NS), spec[key], tokens)
    anchor = tree.find("p:grpSpPr", NS)
    for index, item in enumerate(spec.get("decor", []), 1):
        anchor.addnext(_decor(item, index, tokens)); anchor = anchor.getnext()


def build_preset_reference(tokens, destination=None, base_source=ROOT / RVLAB_ORIGINAL):
    """Build a reference PPTX for any 16:9 or 4:3 preset or custom registered template."""
    destination = Path(destination or (ROOT / tokens["reference"]))
    if tokens.get("preset_id", "rvlab") == "rvlab" and tokens.get("source", {}).get("path") == RVLAB_ORIGINAL:
        return build_reference(ROOT / tokens["source"]["path"], destination, tokens)
    base_source = Path(base_source)
    width_in = float(tokens["slide"]["width_in"])
    height_in = float(tokens["slide"]["height_in"])
    sx, sy = width_in / 10.0, height_in / 7.5
    with ZipFile(base_source) as z:
        parts = {
            n: z.read(n)
            for n in z.namelist()
            if not n.startswith(("ppt/slides/", "ppt/notesSlides/", "ppt/media/")) and n != "ppt/metadata"
        }
    presentation = parse(parts["ppt/presentation.xml"])
    for e in presentation.findall("p:sldIdLst", NS):
        presentation.remove(e)
    sld_sz = presentation.find("p:sldSz", NS)
    if sld_sz is not None:
        sld_sz.set("cx", str(round(width_in * EMU)))
        sld_sz.set("cy", str(round(height_in * EMU)))
        if tokens["slide"].get("aspect") == "16:9":
            sld_sz.set("type", "screen16x9")
        elif tokens["slide"].get("aspect") == "4:3":
            sld_sz.set("type", "screen4x3")
        else:
            sld_sz.attrib.pop("type", None)
    parts["ppt/presentation.xml"] = xml(presentation)
    relationships = parse(parts["ppt/_rels/presentation.xml.rels"])
    for rel in list(relationships):
        if rel.get("Type", "").endswith(("/slide", "/presentationmetadata")):
            relationships.remove(rel)
    parts["ppt/_rels/presentation.xml.rels"] = xml(relationships)
    parts["docProps/core.xml"] = xml(etree.Element("{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}coreProperties", nsmap={"cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}))
    for name in list(parts):
        if name.endswith(".rels") and name.startswith(("ppt/slideLayouts/_rels/", "ppt/slideMasters/_rels/")):
            rels_root = parse(parts[name])
            for rel in list(rels_root):
                if rel.get("Type", "").endswith("/image"):
                    rels_root.remove(rel)
            parts[name] = xml(rels_root)
    master_path = "ppt/slideMasters/slideMaster1.xml"
    master = parse(parts[master_path])
    _transform_xml_for_preset(master, tokens, sx, sy, strip_pics=True)
    if "style" in tokens:
        set_background(master, _color(tokens["style"].get("background", "background"), tokens))
    for theme_name in [n for n in parts if n.startswith("ppt/theme/") and n.endswith(".xml")]:
        theme_root = parse(parts[theme_name])
        _transform_xml_for_preset(theme_root, tokens, 1.0, 1.0, strip_pics=False)
        parts[theme_name] = xml(theme_root)
    ids = master.find("p:sldLayoutIdLst", NS)
    next_id = max(int(i.get("id")) for i in ids) + 1
    relpath = "ppt/slideMasters/_rels/slideMaster1.xml.rels"
    master_rels = parse(parts[relpath])
    existing_ids = {r.get("Id") for r in master_rels}
    originals = {}
    for i in range(1, 5):
        layout_name = f"ppt/slideLayouts/slideLayout{i}.xml"
        layout_root = parse(parts[layout_name])
        _transform_xml_for_preset(layout_root, tokens, sx, sy, strip_pics=True)
        if "style" in tokens:
            apply_layout_style(layout_root, {1: "cover", 2: "toc", 3: "section", 4: "content"}[i], tokens)
        parts[layout_name] = xml(layout_root)
        originals[i] = layout_root
    derivatives = []
    for src, name, typ in [(1, "Title Slide", "title"), (3, "Section Header", "secHead")]:
        layout = deepcopy(originals[src]); layout.set("type", typ); layout.set("preserve", "1")
        layout.find("p:cSld", NS).set("name", name)
        for shape in layout.findall("p:cSld/p:spTree/p:sp", NS):
            ph = placeholder(shape)
            if ph is not None and ph.get("type") == "ctrTitle" and src == 3:
                ph.set("type", "title")
        derivatives.append((src, layout))
    for name, count in [("Title and Content", 1), ("Two Content", 2), ("Comparison", 2), ("Content with Caption", 1), ("Blank", 1), ("RVLAB Figure Grid 3", 3), ("RVLAB Figure Grid 4", 4)]:
        derivatives.append((4, content_layout(originals[4], name, count, tokens)))
    types = parse(parts["[Content_Types].xml"])
    for idx, (src, layout) in enumerate(derivatives, 5):
        name = f"ppt/slideLayouts/slideLayout{idx}.xml"
        parts[name] = xml(layout)
        parts[f"ppt/slideLayouts/_rels/slideLayout{idx}.xml.rels"] = parts[f"ppt/slideLayouts/_rels/slideLayout{src}.xml.rels"]
        rid = f"rvlabLayout{idx}"
        if rid in existing_ids:
            raise ValueError("Duplicate relationship ID")
        etree.SubElement(master_rels, f"{{{NS['rel']}}}Relationship", Id=rid, Type=NS["r"] + "/slideLayout", Target=f"../slideLayouts/slideLayout{idx}.xml")
        etree.SubElement(ids, f"{{{NS['p']}}}sldLayoutId", id=str(next_id + idx - 5), attrib={f"{{{NS['r']}}}id": rid})
        etree.SubElement(types, f"{{{NS['ct']}}}Override", PartName="/" + name, ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml")
    for t in list(types):
        if t.get("PartName", "").lstrip("/") not in parts and t.get("PartName"):
            types.remove(t)
    parts[master_path] = xml(master)
    parts[relpath] = xml(master_rels)
    parts["[Content_Types].xml"] = xml(types)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=destination.parent, suffix=".pptx", delete=False) as f:
        temporary = Path(f.name)
    try:
        with ZipFile(temporary, "w", ZIP_DEFLATED) as z:
            for name, data in sorted(parts.items()):
                info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
                info.compress_type = ZIP_DEFLATED
                z.writestr(info, data)
        validate_package(temporary)
        prs = Presentation(temporary)
        if len(prs.slides) != 0 or len(prs.slide_masters) != 1 or len(prs.slide_layouts) != 13:
            raise ValueError("Invalid derived reference structure")
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)
    return {"preset_id": tokens.get("preset_id"), "reference": str(destination.relative_to(ROOT) if destination.is_relative_to(ROOT) else destination), "slide_inches": [width_in, height_in]}


SPECIMEN = {
    "cover": {0: "발표 제목이 들어가는 자리", 1: "부제 · Subtitle", 2: "발표자 (Presenter)", 3: "소속 첫 줄\n소속 둘째 줄", 4: "발표일시 : 2026.01.01"},
    "toc": [(0, "Chapter 1: 배경과 문제"), (1, "토픽 A"), (1, "토픽 B"), (0, "Chapter 2: 방법"), (1, "토픽 C"), (0, "Chapter 3: 결과와 다음 단계")],
    "section": {0: "섹션 제목", 1: "Section subtitle"},
    "content": [(0, "핵심 문장: 이 슬라이드의 결론 한 줄"), (1, "근거가 되는 첫 번째 불릿"), (1, "두 번째 불릿"), (2, "필요하면 한 단계 더 들여 쓴 설명")],
    "columns": [[(0, "왼쪽 단"), (1, "관찰 A"), (1, "관찰 B")], [(0, "오른쪽 단"), (1, "제안 A"), (1, "제안 B")]],
}


def _fill(frame, lines):
    for i, (level, value) in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text, paragraph.level = value, level


def write_specimen(tokens, root=ROOT):
    """templates/<id>/preview.pdf and preview.png: one sample slide per layout, so the theme can be seen without building a deck.

    The reference itself must stay slide-free because the build adds slides onto it."""
    import pymupdf
    from .preview import contact_sheet, render_pdf
    reference = Path(root) / tokens["reference"]
    folder = reference.parent
    prs = Presentation(reference)
    layouts = {layout.name: layout for layout in prs.slide_layouts}
    names = tokens["layouts"]

    def slide(layout):
        new = prs.slides.add_slide(layouts[names[layout]])
        return new, {ph.placeholder_format.idx: ph for ph in new.placeholders}

    _, phs = slide("cover")
    for idx, value in SPECIMEN["cover"].items():
        if idx in phs:
            phs[idx].text_frame.text = value
    _, phs = slide("toc")
    _fill(phs[1].text_frame, SPECIMEN["toc"])
    _, phs = slide("section")
    for idx, value in SPECIMEN["section"].items():
        phs[idx].text_frame.text = value
    _, phs = slide("content")
    phs[0].text_frame.text = "토픽 이름"
    _fill(phs[1].text_frame, SPECIMEN["content"])
    _, phs = slide("columns")
    phs[0].text_frame.text = "2단 비교"
    for idx, lines in enumerate(SPECIMEN["columns"], 1):
        _fill(phs[idx].text_frame, lines)
    _, phs = slide("figure_3")
    phs[0].text_frame.text = "그림 그리드"
    for idx in (1, 2, 3):
        phs[idx].text_frame.text = f"그림 {idx} 자리"
    closing, phs = slide("section")
    phs[0].text_frame.text = tokens["labels"]["ko"]["closing"]
    phs[0].left, phs[0].top, phs[0].width, phs[0].height = [round(v * EMU) for v in tokens["geometry_in"]["closing_title"]]
    closing.shapes._spTree.remove(phs[1]._element)
    with tempfile.TemporaryDirectory() as directory:
        deck = Path(directory) / "preview.pptx"
        prs.save(deck)
        render_pdf(deck, folder / "preview.pdf")
        images = []
        with pymupdf.open(folder / "preview.pdf") as pdf:
            for number, page in enumerate(pdf, 1):
                path = Path(directory) / f"slide-{number}.png"
                page.get_pixmap(dpi=60, alpha=False).save(path)
                images.append(path)
        contact_sheet(images, ["Cover", "Contents", "Section", "Content", "Two columns", "Figure grid (3)", "Closing"], folder / "preview.png")
    return {"pdf": str(folder / "preview.pdf"), "png": str(folder / "preview.png")}


def ensure_preset_reference(tokens):
    ref_path = ROOT / tokens["reference"]
    if tokens.get("preset_id", "rvlab") == "rvlab":
        if not ref_path.is_file():
            raise ValueError("Reference missing: uv run python -m rvppt.template")
    else:
        # Derived presets regenerate byte-identically in milliseconds, so every build picks up DESIGN.md edits.
        build_preset_reference(tokens, ref_path)
    return ref_path


def register_template(source_pptx, preset_id, *, name=None, description=None, root=ROOT):
    """Inspect a user-supplied .pptx file and register it as a reusable design preset."""
    if not PRESET_ID_PATTERN.fullmatch(str(preset_id)):
        raise ValueError("Preset ID must be lowercase alphanumeric/hyphens")
    existing = design_path(preset_id, root)
    if existing.is_file() and (read_design(existing) or {}).get("mode") != "custom_template":
        raise ValueError(f"'{preset_id}' is a built-in template; choose another ID")
    source_path = Path(source_pptx).resolve()
    if not source_path.is_file() or source_path.suffix.lower() != ".pptx":
        raise ValueError(f"Template must be an existing .pptx file: {source_pptx}")
    manifest = extract_manifest(source_path)
    w_in = round(manifest["slide_inches"][0], 3)
    h_in = round(manifest["slide_inches"][1], 3)
    ratio = w_in / h_in
    aspect = "16:9" if abs(ratio - 16 / 9) < 0.05 else "4:3" if abs(ratio - 4 / 3) < 0.05 else f"{w_in}:{h_in}"
    custom_dir = Path(root) / "templates" / preset_id
    custom_dir.mkdir(parents=True, exist_ok=True)
    stored_pptx = custom_dir / "source.pptx"
    if source_path != stored_pptx.resolve():
        shutil.copy2(source_path, stored_pptx)
    manifest["source"] = {"path": f"templates/{preset_id}/source.pptx", "sha256": sha256(stored_pptx)}
    (custom_dir / "template-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")

    # Extract theme accent colors if present, falling back to clean defaults
    theme_colors = []
    for theme in manifest.get("themes", {}).values():
        for key in ("accent1", "dk2", "accent2", "accent3"):
            for item in theme.get(key, []):
                val = item.get("val", "")
                if re.fullmatch(r"[0-9A-Fa-f]{6}", val) and val.upper() not in ("FFFFFF", "000000"):
                    theme_colors.append(val.upper())
    for c in manifest.get("layout_direct_colors", []):
        if re.fullmatch(r"[0-9A-Fa-f]{6}", c) and c.upper() not in ("FFFFFF", "000000") and c.upper() not in theme_colors:
            theme_colors.append(c.upper())
    primary = theme_colors[0] if len(theme_colors) >= 1 else "1D4ED8"
    band = theme_colors[1] if len(theme_colors) >= 2 else primary
    secondary = theme_colors[2] if len(theme_colors) >= 3 else "60A5FA"

    # 16:9 inherits modern-minimal's structure; other ratios start from the RVLAB structure.
    base_tokens = load_tokens(preset="modern-minimal" if aspect == "16:9" else "rvlab", root=root)
    sx, sy = w_in / base_tokens["slide"]["width_in"], h_in / base_tokens["slide"]["height_in"]
    scale = lambda box: [round(v * f, 2) for v, f in zip(box, (sx, sy, sx, sy))]
    tokens = deepcopy(base_tokens)
    tokens.pop("source", None)  # The RVLAB original belongs to the rvlab template only.
    tokens.update({
        "preset_id": preset_id,
        "name": name or f"Custom Template ({preset_id})",
        "description": description or f"Registered from {source_path.name} ({aspect}, {w_in}x{h_in} in)",
        "mode": "custom_template",
        "custom_source": manifest["source"],
        "reference": f"templates/{preset_id}/reference.pptx",
        "slide": {"width_in": w_in, "height_in": h_in, "aspect": aspect, "masters": 1, "original_layouts": 4},
        "palette": {
            "primary": primary,
            "band": band,
            "secondary": secondary,
            "section_accent": "FACC15",
            "text": "0F172A",
            "background": "FFFFFF",
        },
        "geometry_in": {
            **{key: scale(base_tokens["geometry_in"][key]) for key in ("title", "header_rule", "body", "closing_title")},
            "content_indent": 0.25,
            "gap": base_tokens["geometry_in"]["gap"],
            "caption_height": base_tokens["geometry_in"]["caption_height"],
        },
    })
    if "style" in tokens:
        # The inherited structure is drawn in the base slide's inches; move every box onto the registered slide.
        def rescale(node):
            if isinstance(node, dict):
                return {k: scale(v) if k == "box" else rescale(v) for k, v in node.items()}
            return [rescale(v) for v in node] if isinstance(node, list) else node
        tokens["style"] = rescale(tokens["style"])
    design = design_path(preset_id, root)
    design.parent.mkdir(parents=True, exist_ok=True)
    # Re-registering refreshes the tokens but keeps conventions someone already wrote below them.
    prose = (f"## 성격과 판단 기준\n\n등록한 `{manifest['source']['path']}`를 열어 보고 성격·어울리는 것·어울리지 않는 것·그림 인상·빌드 후 볼 것을 여기에 적는다. "
             "형식은 `templates/modern-minimal/DESIGN.md`를 참고한다. 구조·장식은 YAML의 `style` 블록에서 바꾼다.\n")
    if design.is_file() and "```yaml\n" in (old := design.read_text(encoding="utf-8")):
        prose = old.split("```yaml\n", 1)[1].split("```\n", 1)[1].lstrip("\n") or prose
    design.write_text(f"# {tokens['name']} 디자인 문서\n\n```yaml\n{yaml.safe_dump(tokens, allow_unicode=True, sort_keys=False)}```\n\n{prose}", encoding="utf-8", newline="\n")
    ref_path = Path(root) / tokens["reference"]
    build_preset_reference(tokens, ref_path, base_source=Path(root) / RVLAB_ORIGINAL)
    preview = write_specimen(tokens, root)
    return {
        "preset_id": preset_id,
        "design": str(design),
        "preview": preview["png"],
        "reference": str(ref_path),
        "manifest": str(custom_dir / "template-manifest.json"),
        "slide_inches": [w_in, h_in],
        "aspect": aspect,
        "palette": tokens["palette"],
    }


def build_reference(source, destination, tokens):
    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve():
        raise ValueError("The original template cannot be overwritten")
    if sha256(source) != tokens["source"]["sha256"]:
        raise ValueError("Original template hash differs from DESIGN.md")
    manifest = extract_manifest(source)
    if manifest["slide_inches"] != [10, 7.5] or manifest["masters"] != 1 or len(manifest["layouts"]) != 4:
        raise ValueError("Unexpected original template structure")
    with ZipFile(source) as z:
        parts = {n: z.read(n) for n in z.namelist() if not n.startswith(("ppt/slides/", "ppt/notesSlides/")) and n != "ppt/metadata"}
    presentation = parse(parts["ppt/presentation.xml"])
    for e in presentation.findall("p:sldIdLst", NS):
        presentation.remove(e)
    parts["ppt/presentation.xml"] = xml(presentation)
    relationships = parse(parts["ppt/_rels/presentation.xml.rels"])
    for rel in list(relationships):
        if rel.get("Type", "").endswith(("/slide", "/presentationmetadata")):
            relationships.remove(rel)
    parts["ppt/_rels/presentation.xml.rels"] = xml(relationships)
    parts["docProps/core.xml"] = xml(etree.Element("{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}coreProperties", nsmap={"cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}))
    types = parse(parts["[Content_Types].xml"])
    for t in list(types):
        if t.get("PartName", "").lstrip("/") not in parts and t.get("PartName"):
            types.remove(t)
    master_path = "ppt/slideMasters/slideMaster1.xml"
    master = parse(parts[master_path])
    ids = master.find("p:sldLayoutIdLst", NS)
    next_id = max(int(i.get("id")) for i in ids) + 1
    relpath = "ppt/slideMasters/_rels/slideMaster1.xml.rels"
    master_rels = parse(parts[relpath])
    existing_ids = {r.get("Id") for r in master_rels}
    originals = {i: parse(parts[f"ppt/slideLayouts/slideLayout{i}.xml"]) for i in range(1, 5)}
    derivatives = []
    for src, name, typ in [(1, "Title Slide", "title"), (3, "Section Header", "secHead")]:
        layout = deepcopy(originals[src]); layout.set("type", typ); layout.set("preserve", "1")
        layout.find("p:cSld", NS).set("name", name)
        for shape in layout.findall("p:cSld/p:spTree/p:sp", NS):
            ph = placeholder(shape)
            if ph is not None and ph.get("type") == "ctrTitle" and src == 3:
                ph.set("type", "title")
        derivatives.append((src, layout))
    for name, count in [("Title and Content", 1), ("Two Content", 2), ("Comparison", 2), ("Content with Caption", 1), ("Blank", 1), ("RVLAB Figure Grid 3", 3), ("RVLAB Figure Grid 4", 4)]:
        derivatives.append((4, content_layout(originals[4], name, count, tokens)))
    for idx, (src, layout) in enumerate(derivatives, 5):
        name = f"ppt/slideLayouts/slideLayout{idx}.xml"
        parts[name] = xml(layout)
        parts[f"ppt/slideLayouts/_rels/slideLayout{idx}.xml.rels"] = parts[f"ppt/slideLayouts/_rels/slideLayout{src}.xml.rels"]
        rid = f"rvlabLayout{idx}"
        if rid in existing_ids:
            raise ValueError("Duplicate relationship ID")
        etree.SubElement(master_rels, f"{{{NS['rel']}}}Relationship", Id=rid, Type=NS["r"] + "/slideLayout", Target=f"../slideLayouts/slideLayout{idx}.xml")
        etree.SubElement(ids, f"{{{NS['p']}}}sldLayoutId", id=str(next_id + idx - 5), attrib={f"{{{NS['r']}}}id": rid})
        etree.SubElement(types, f"{{{NS['ct']}}}Override", PartName="/" + name, ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml")
    parts[master_path] = xml(master)
    parts[relpath] = xml(master_rels)
    parts["[Content_Types].xml"] = xml(types)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=destination.parent, suffix=".pptx", delete=False) as f:
        temporary = Path(f.name)
    try:
        with ZipFile(temporary, "w", ZIP_DEFLATED) as z:
            for name, data in sorted(parts.items()):
                # Stable zip timestamps make regeneration byte reproducible.
                info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
                info.compress_type = ZIP_DEFLATED
                z.writestr(info, data)
        validate_package(temporary)
        prs = Presentation(temporary)
        if len(prs.slides) != 0 or len(prs.slide_masters) != 1 or len(prs.slide_layouts) != 13:
            raise ValueError("Invalid derived reference structure")
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--preset", help="Optional preset ID to generate reference for (default: all presets)")
    args = parser.parse_args()
    tokens = load_tokens()
    manifest = build_reference(ROOT / tokens["source"]["path"], ROOT / tokens["reference"], tokens)
    target = ROOT / "templates/rvlab/template-manifest.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print(f"Reference generated: {tokens['reference']} (1 master, 13 layouts, 0 specimen slides)")
    print(f"Original preserved: {manifest['source']['sha256']}")
    if not args.preset or args.preset == "rvlab":
        print(f"Preview: {write_specimen(tokens)['png']}")
    for item in list_presets()["presets"]:
        pid = item["preset_id"]
        if pid == "rvlab":
            continue
        if args.preset and args.preset != pid:
            continue
        ptokens = load_tokens(preset=pid)
        build_preset_reference(ptokens)
        print(f"Preset reference generated: {ptokens['reference']} ({ptokens['slide']['aspect']}), preview {write_specimen(ptokens)['png']}")


if __name__ == "__main__":
    main()
