import json
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import etree
from pptx import Presentation

from rvppt.doctor import skill_integrity
from rvppt.graph import graph_check, ignored
from rvppt.preview import fit_picture
from rvppt.template import ROOT, NS, EMU, build_reference, load_tokens, parse, sha256, slots, validate_package


class TemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tokens = load_tokens()
        cls.original = ROOT / cls.tokens["source"]["path"]
        cls.temp = tempfile.TemporaryDirectory()
        cls.reference = Path(cls.temp.name) / "reference.pptx"
        cls.manifest = build_reference(cls.original, cls.reference, cls.tokens)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_original_and_original_visual_parts_preserved(self):
        self.assertEqual(sha256(self.original), self.tokens["source"]["sha256"])
        with ZipFile(self.original) as source, ZipFile(self.reference) as derived:
            for name in source.namelist():
                if name.startswith(("ppt/slideLayouts/", "ppt/media/", "ppt/theme/")):
                    self.assertEqual(source.read(name), derived.read(name), name)
            master = "ppt/slideMasters/slideMaster1.xml"
            left, right = parse(source.read(master)), parse(derived.read(master))
            for root in (left, right):
                root.remove(root.find("p:sldLayoutIdLst", NS))
            self.assertEqual(etree.tostring(left, method="c14n"), etree.tostring(right, method="c14n"))

    def test_reference_contains_no_old_slides_and_required_layouts(self):
        prs = Presentation(self.reference)
        self.assertEqual(len(prs.slides), 0)
        self.assertEqual(len(prs.slide_masters), 1)
        self.assertEqual((prs.slide_width, prs.slide_height), (10 * EMU, 7.5 * EMU))
        names = {layout.name for layout in prs.slide_layouts}
        self.assertTrue({"Title Slide", "Section Header", "Title and Content", "Two Content", "Comparison", "Content with Caption", "Blank", "RVLAB Figure Grid 3", "RVLAB Figure Grid 4"}.issubset(names))
        with ZipFile(self.reference) as z:
            self.assertFalse(any(n.startswith(("ppt/slides/", "ppt/notesSlides/")) for n in z.namelist()))
            self.assertNotIn("ppt/metadata", z.namelist())

    def test_top_bar_keeps_original_rule_and_body_stays_below_it(self):
        prs = Presentation(self.reference)
        original = {l.name: l for l in Presentation(self.original).slide_layouts}["OBJECT"]
        original_rule = next(s for s in original.shapes if s.shape_id == 41)
        for layout in prs.slide_layouts:
            if layout.name not in {"Title and Content", "Two Content", "Comparison", "Content with Caption", "RVLAB Figure Grid 3", "RVLAB Figure Grid 4"}:
                continue
            title = next(p for p in layout.placeholders if p.placeholder_format.idx == 0)
            rule = next(s for s in layout.shapes if s.shape_id == 41)
            objects = [p for p in layout.placeholders if 1 <= p.placeholder_format.idx <= 4]
            # RVLAB's top bar: a white-filled short title tab masking the rule, which runs on to the right.
            self.assertEqual([round(v / EMU, 2) for v in (title.left, title.top, title.width, title.height)], self.tokens["geometry_in"]["title"])
            self.assertIsNotNone(title._element.find("p:spPr/a:solidFill", NS))
            self.assertEqual((rule.left, rule.top, rule.width, rule.height), (original_rule.left, original_rule.top, original_rule.width, original_rule.height))
            for shape in objects:
                self.assertGreaterEqual(shape.top, title.top + title.height)
                self.assertLess(shape.top + shape.height, 7.23 * EMU)
                ph = shape._element.find("p:nvSpPr/p:nvPr/p:ph", NS)
                self.assertNotIn("type", ph.attrib, "Pandoc object placeholder must omit type")
                for other in objects:
                    if shape is not other:
                        overlaps = shape.left < other.left + other.width and other.left < shape.left + shape.width and shape.top < other.top + other.height and other.top < shape.top + shape.height
                        self.assertFalse(overlaps)

    def test_internal_relationships_and_shape_ids_are_valid(self):
        self.assertTrue(validate_package(self.reference))
        with ZipFile(self.reference) as z:
            for name in z.namelist():
                if name.startswith("ppt/slideLayouts/") and name.endswith(".xml"):
                    ids = parse(z.read(name)).xpath(".//p:cNvPr/@id", namespaces=NS)
                    self.assertEqual(len(ids), len(set(ids)), name)
            root = parse(z.read("ppt/slideMasters/slideMaster1.xml"))
            rels = parse(z.read("ppt/slideMasters/_rels/slideMaster1.xml.rels"))
            known = {r.get("Id") for r in rels}
            self.assertTrue(set(root.xpath(".//@r:id", namespaces=NS)).issubset(known))

    def test_bad_original_preserves_last_success_and_cannot_overwrite_original(self):
        original_hash, previous = sha256(self.original), self.reference.read_bytes()
        bad = dict(self.tokens, source={"sha256": "wrong"})
        with self.assertRaises(ValueError):
            build_reference(self.original, self.reference, bad)
        self.assertEqual(self.reference.read_bytes(), previous)
        with self.assertRaises(ValueError):
            build_reference(self.original, self.original, self.tokens)
        self.assertEqual(sha256(self.original), original_hash)

    def test_generation_is_byte_reproducible(self):
        other = Path(self.temp.name) / "again.pptx"
        build_reference(self.original, other, self.tokens)
        self.assertEqual(sha256(self.reference), sha256(other))

    def test_broken_relationship_is_rejected(self):
        broken = Path(self.temp.name) / "broken.pptx"
        with ZipFile(self.reference) as source, ZipFile(broken, "w", ZIP_DEFLATED) as target:
            for name in source.namelist():
                if name != "ppt/media/image1.png":
                    target.writestr(name, source.read(name))
        with self.assertRaisesRegex(ValueError, "relationship target"):
            validate_package(broken)

    def test_picture_fit_preserves_aspect_and_stays_in_slots(self):
        for count in (1, 2, 3, 4):
            for box in slots(self.tokens, count):
                for size in [(1600, 1000), (1000, 1600), (1600, 1600), (1800, 800)]:
                    x, y, w, h = fit_picture(size, box)
                    self.assertAlmostEqual(w / h, size[0] / size[1])
                    self.assertGreaterEqual(x, box[0])
                    self.assertGreaterEqual(y, box[1])
                    self.assertLessEqual(x + w, box[0] + box[2] + 1e-9)
                    self.assertLessEqual(y + h, box[1] + box[3] + 1e-9)
        with self.assertRaises(ValueError):
            slots(self.tokens, 5)


class SkillIntegrityTests(unittest.TestCase):
    def test_vendored_skills_match_manifest_and_tampering_is_reported(self):
        self.assertEqual(skill_integrity(), [])
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            shutil.copytree(ROOT / "skills", root / "skills", symlinks=True)
            (root / "skills/im-not-ai/LICENSE").write_text("changed")
            (root / "skills/pdf/SKILL.md").unlink()
            problems = skill_integrity(root)
            self.assertTrue(any("im-not-ai/LICENSE: modified" in p for p in problems))
            self.assertTrue(any("pdf/SKILL.md: missing" in p for p in problems))

    def test_symlinks_checked_out_as_plain_files_are_reported(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            shutil.copytree(ROOT / "skills", root / "skills", symlinks=True)
            vendored = root / "skills/im-not-ai/codex/skills/humanize-korean/references"
            vendored.unlink()
            vendored.write_bytes(b"../../../skills/humanize-korean/references")
            (root / ".claude/skills").mkdir(parents=True)
            (root / ".claude/skills/rvlab-ppt").write_bytes(b"../../skills/rvlab-ppt")
            problems = skill_integrity(root)
            for path in ("skills/im-not-ai/codex/skills/humanize-korean/references", ".claude/skills/rvlab-ppt"):
                self.assertTrue(any(p.startswith(f"{path}: symlink checked out as a plain file") for p in problems), path)


class GraphCheckTests(unittest.TestCase):
    def test_ignore_rules_follow_graphifyignore_semantics(self):
        patterns = ["/tasks/", "/docs/*.png", "node_modules/", "uv.lock"]
        for path, expected in {"tasks/a/slides.md": True, "docs/x.png": True, "docs/design/x.png": False,
                               "web/node_modules/a.js": True, "uv.lock": True, "tools/rvppt/tasks.py": False, "docs/tasks.md": False}.items():
            self.assertEqual(ignored(path, patterns), expected, path)

    def test_snapshot_detects_changes_new_inputs_and_task_leaks(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            subprocess.run(["git", "init", "-q"], cwd=root, check=True)
            (root / ".graphifyignore").write_text("/tasks/\n/graphify-out/\n")
            (root / "tasks/t").mkdir(parents=True); (root / "tasks/t/slides.md").write_text("private deck")
            (root / "a.py").write_text("x = 1\n")
            out = root / "graphify-out"; out.mkdir()
            graph = {"nodes": [{"id": "a", "source_file": "a.py"}], "links": []}
            (out / "graph.json").write_text(json.dumps(graph))
            for name in ("GRAPH_REPORT.md", "graph.html"):
                (out / name).write_text("ok")
            self.assertTrue(graph_check(root, record=True)["ok"])
            (root / "a.py").write_text("x = 2\n")
            (root / "b.py").write_text("y = 1\n")
            result = graph_check(root)
            self.assertEqual((result["ok"], result["changed"], result["added"]), (False, ["a.py"], ["b.py"]))
            graph["nodes"].append({"id": "t", "source_file": "tasks/t/slides.md"})
            (out / "graph.json").write_text(json.dumps(graph))
            self.assertEqual(graph_check(root)["leaks"], ["tasks/t/slides.md"])
            with self.assertRaisesRegex(ValueError, "Excluded"):
                graph_check(root, record=True)


class IgnoreTests(unittest.TestCase):
    def test_only_final_example_outputs_are_trackable(self):
        cases = {
            "tasks/ordinary/20261011_발표.pptx": True,
            "tasks/ordinary/20261011_발표.pdf": True,
            "tasks/ordinary/.build/output/slide-01.png": True,
            "tasks/ordinary/.rvppt-approval.tmp": True,
            "tasks/ordinary/plan/.rvppt-x.tmp": True,
            "tasks/ordinary/slides.md": True,
            "tasks/ordinary/sources/draft.pptx": True,
            "tasks/ordinary/sources/paper.pdf": True,
            "tasks/261011-example-il-progress/sources/eval.csv": False,
            "tasks/261011-example-il-progress/sources/new-data.csv": False,
            "tasks/ordinary/figures/curve.png": True,
            "tasks/261011-example-il-progress/20261011_예제.pptx": False,
            "tasks/261011-example-il-progress/20261011_예제.pdf": False,
            "tasks/261011-example-il-progress/.build/output/contact-sheet.png": True,
            "profile.yaml": True,
            "profile.example.yaml": False,
            "docs/design/build/layout-specimens.pptx": True,
            "graphify-out/graph.json": False,
            "graphify-out/snapshot.json": False,
            ".env.local": True,
            "uv.lock": False,
        }
        for path, ignored in cases.items():
            result = subprocess.run(["git", "check-ignore", "-q", "--no-index", path], cwd=ROOT)
            self.assertIn(result.returncode, (0, 1))
            self.assertEqual(result.returncode == 0, ignored, path)


if __name__ == "__main__":
    unittest.main()
