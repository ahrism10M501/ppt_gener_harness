from copy import deepcopy
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
import yaml

from build_fixture import make_task
from rvppt import build, markdown, polish, project
from rvppt.preview import resolve_fonts
from rvppt.state import StateError, STAGES, read_document, status
from rvppt.template import ROOT, load_tokens, sha256


class BuildTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fixture = tempfile.TemporaryDirectory()
        cls.original = make_task(Path(cls.fixture.name) / "261011-build-test")

    @classmethod
    def tearDownClass(cls):
        cls.fixture.cleanup()

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.task = Path(self.temporary.name) / self.original.name
        shutil.copytree(self.original, self.task)
        self.tokens = load_tokens()

    def tearDown(self):
        self.temporary.cleanup()

    def edit(self, stage, transform):
        document = read_document(self.task / STAGES[stage - 1], self.task.name, stage)
        body = transform(document.body)
        document.path.write_text("---\n" + yaml.safe_dump(document.metadata, allow_unicode=True, sort_keys=False) + "---\n" + body, encoding="utf-8")

    def slide(self, slides, identifier):
        return deepcopy(next(s for s in slides if s["id"] == identifier))

    def test_pandoc_parser_builds_rvlab_deck_structure(self):
        _, slides, _ = markdown.parse_task(self.task)
        self.assertEqual([s["layout"] for s in slides], ["cover", "toc", "section", "content", "content", "columns", "columns", "section", "content", "figure_2", "figure_3", "figure_4", "content", "content", "content", "closing", "section", "content"])
        self.assertEqual(slides[0]["date_label"], "발표일시 : 2026.10.11")
        self.assertEqual(slides[1]["chapters"], [("Chapter 1: 기능 검증", ["표와 수식", "2단 본문"]), ("Chapter 2: 그림 배치", ["그림 슬롯", "그림 요청", "Mermaid", "Code"])])
        self.assertEqual((slides[2]["title"], slides[2]["subtitle"]), ("기능 검증", "Feature QA"))
        self.assertEqual((slides[3]["topic"], slides[3]["headline"]), ("표와 수식", "합성 표의 모든 셀을 편집할 수 있다"))
        self.assertEqual(slides[15]["title"], "감사합니다.")
        self.assertEqual(len(self.slide(slides, "S08")["slots"]), 4)
        self.assertEqual(sum(n["t"] == "Math" for n in markdown.nodes(self.slide(slides, "S02")["slots"])), 4)
        self.assertTrue(all(s["notes"] for s in slides if s["layout"] not in ("cover", "toc", "section", "closing")))

    def test_top_bar_topic_must_be_a_short_name(self):
        self.edit(4, lambda s: s.replace("## 2단 본문", "## Where is Skip-Connection Now"))
        with self.assertRaisesRegex(StateError, "Top-bar topic"):
            build.lint(self.task)

    def test_content_requires_a_headline_slide(self):
        self.edit(4, lambda s: s.replace("## 표와 수식\n", "## 표와 수식\n\n- 핵심 문장 없는 내용\n", 1))
        with self.assertRaisesRegex(StateError, "headline"):
            build.lint(self.task)

    def test_missing_figure_is_requested_without_failing(self):
        result = build.lint(self.task)
        self.assertEqual(result["missing_figures"], [{"slide_id": "S09", "path": "figures/requested-curve.png", "request": "x축 epoch, y축 성공률인 합성 곡선. sources/synthetic.csv로 작성"}])
        with self.assertRaisesRegex(StateError, "Missing task image"):
            markdown.asset_path(self.task, "sources/missing.png")

    def test_output_uses_lab_file_naming(self):
        self.assertEqual(build.output_stem({"date": "2026-08-31", "title": "학습루프개선 및 실험"}), "20260831_학습루프개선_및_실험")

    def test_new_project_fills_cover_from_profile(self):
        profile = self.task.parent / "profile.yaml"
        with self.assertRaisesRegex(StateError, "profile.yaml is missing"):
            project.new_project("lab-meeting", "학습루프 개선", profile_path=profile, tasks=self.task.parent)
        profile.write_text("name: 홍길동\nname_en: Gildong Hong\naffiliation: [경기대학교 전자공학과 N학년, Robotics and Vision Lab Robot Team]\n", encoding="utf-8")
        created = Path(project.new_project("lab-meeting", "학습루프 개선", "Train Loop", "2026-08-31", profile_path=profile, tasks=self.task.parent)["task"])
        self.assertEqual(created.name, "260831-lab-meeting")
        self.assertTrue(all((created / d).is_dir() for d in ("figures", "sources", "plan")))
        _, slides, _ = markdown.parse_task(created)
        self.assertEqual((slides[0]["presenter"], slides[0]["date_label"]), ("홍길동 (Gildong Hong)", "발표일시 : 2026.08.31"))
        self.assertEqual(slides[1]["layout"], "toc")
        self.assertEqual(len(build.lint(created)["missing_figures"]), 1)
        compact = Path(project.new_project("compact-date", "날짜 형식", day="20261012", profile_path=profile, tasks=self.task.parent)["task"])
        self.assertEqual(compact.name, "261012-compact-date")
        document = (compact / "slides.md").read_text(encoding="utf-8")
        (compact / "slides.md").write_text(document.replace("date: '2026-10-12'", "date: '20261012'"), encoding="utf-8")
        self.assertEqual(markdown.parse_task(compact)[1][0]["date_label"], "발표일시 : 2026.10.12")
        self.assertEqual([f["slide_id"] for f in build.lint(compact)["findings"]], ["cover"])  # created without --subtitle
        self.assertEqual(build.lint(created)["findings"], [])
        with self.assertRaisesRegex(StateError, "already exists"):
            project.new_project("lab-meeting", "중복", day="2026-08-31", profile_path=profile, tasks=self.task.parent)

    def test_polish_check_accepts_wording_but_not_fact_or_structure_changes(self):
        before = self.task / ".build/editorial/before.md"; before.parent.mkdir(parents=True)
        shutil.copy2(self.task / "slides.md", before)
        original = (self.task / "slides.md").read_text(encoding="utf-8")
        self.edit(4, lambda s: s.replace("- 비교 기준을 기록한다", "- 비교 기준을 남긴다"))
        result = polish.polish_check(self.task, before)
        self.assertTrue(result["ok"], result)
        cases = {
            "Numbers changed": ("| B | 18 | 정렬 |", "| B | 19 | 정렬 |"),
            "structure changed": ("figures/chart.svg)", "figures/chart-new.svg)"),
            "structure changed ": (r"\frac{1}{N}", r"\frac{2}{N}"),
            "structure changed  ": ("### 합성 표의 모든 셀을 편집할 수 있다", "### 표의 셀을 모두 편집한다"),
            "Frontmatter": ("RVLAB 제작·QA 검증", "RVLAB 검증"),
        }
        for message, (old, new) in cases.items():
            with self.subTest(message=message):
                (self.task / "slides.md").write_text(original.replace(old, new, 1), encoding="utf-8")
                result = polish.polish_check(self.task, before)
                self.assertFalse(result["ok"])
                self.assertIn(message.strip(), " ".join(result["errors"]))

    def test_smart_quotes_remain_in_measured_and_expected_text(self):
        ast = json.loads(markdown.pandoc(["-f", "markdown", "-t", "json"], '이 문장의 "범위"를 보존한다.'))
        paragraph = ast["blocks"][0]
        self.assertEqual(markdown.plain(paragraph).strip(), "이 문장의 “범위”를 보존한다.")
        self.assertEqual("".join(v for kind, v in build.inline_parts(paragraph["c"])), "이 문장의 “범위”를 보존한다.")

    def test_citation_display_text_ignores_pandoc_metadata(self):
        ast = json.loads(markdown.pandoc(["-f", "markdown", "-t", "json"], "Observed [@synthetic2026]."))
        paragraph = ast["blocks"][0]
        self.assertEqual(markdown.plain(paragraph).strip(), "Observed [@synthetic2026].")
        self.assertEqual("".join(v for _, v in build.inline_parts(paragraph["c"])), "Observed [@synthetic2026].")

    def test_layout_attributes_are_rejected(self):
        self.edit(4, lambda s: s.replace("figures/chart.svg)", "figures/chart.svg){width=99%}"))
        with self.assertRaisesRegex(StateError, "attributes"):
            build.lint(self.task)

    def test_paths_and_remote_images_are_rejected(self):
        for path in ("../figures/chart.svg", "figures/../../private.png", "/private.png", "https://example.com/p.png", "assets/wide.png", "figures/notes.txt"):
            with self.subTest(path=path), self.assertRaises(StateError):
                markdown.asset_path(self.task, path)

    def test_symlink_images_are_rejected(self):
        link = self.task / "figures/link.png"
        link.symlink_to(self.task / "figures/wide.png")
        with self.assertRaisesRegex(StateError, "symlink"):
            markdown.asset_path(self.task, "figures/link.png")

    def test_excess_bullets_and_depth_are_rejected(self):
        self.edit(4, lambda s: s.replace("- 실제 근거를 추가한다", "- 실제 근거를 추가한다\n- 항목 C\n- 항목 D"))
        with self.assertRaisesRegex(StateError, "5 bullets"):
            build.lint(self.task)

    def test_three_levels_are_rejected(self):
        self.edit(4, lambda s: s.replace("  - 확인할 조건", "  - 확인할 조건\n    - 깊이 3"))
        with self.assertRaisesRegex(StateError, "2 levels"):
            build.lint(self.task)

    def test_table_row_limit(self):
        self.edit(4, lambda s: s.replace("| C | 24 | 폰트 |", "\n".join(["| C | 24 | 폰트 |"] * 7)))
        with self.assertRaisesRegex(StateError, "8 rows"):
            build.lint(self.task)

    def test_columns_are_required(self):
        self.edit(4, lambda s: s.replace(":::: columns", ":::: arbitrary", 1))
        with self.assertRaisesRegex(StateError, "Allowed blocks"):
            build.lint(self.task)

    def test_mermaid_styling_is_rejected(self):
        self.edit(4, lambda s: s.replace("flowchart LR", "flowchart LR\nstyle A fill:#ff0000"))
        with self.assertRaisesRegex(StateError, "Mermaid styling"):
            build.lint(self.task)

    def test_equation_colors_are_rejected(self):
        self.edit(4, lambda s: s.replace("x^2 + y^2", r"\color{red}{x^2}"))
        with self.assertRaisesRegex(StateError, "LaTeX styling"):
            build.lint(self.task)

    def test_headlines_that_differ_from_outline_are_draft_edits(self):
        self.edit(4, lambda s: s.replace("### 합성 표의 모든 셀을 편집할 수 있다", "### 연구자가 고친 핵심 문장"))
        markdown.parse_task(self.task)  # Researchers can still compile the draft.
        stage = status(self.task)["stages"][3]
        self.assertEqual(stage["effective_status"], "invalid")
        self.assertIn("Slide topics and headlines differ from the outline", stage["issues"])

    def test_headline_quotes_are_not_rewritten(self):
        new = '"합성 기능"을 검증한다'
        self.edit(4, lambda s: s.replace("### 합성 표의 모든 셀을 편집할 수 있다", "### " + new))
        _, slides, _ = markdown.parse_task(self.task)
        self.assertEqual(self.slide(slides, "S01")["headline"], new)

    def test_hand_edited_output_is_never_overwritten(self):
        edited = self.task / "20261011_x.pptx"; edited.write_bytes(b"researcher polish")
        staging = self.task / ".build/staging"; staging.mkdir(parents=True)
        with self.assertRaisesRegex(StateError, "edited by hand"):
            build.publish(staging, self.task, {"20261011_x.pptx": "0" * 64})
        self.assertEqual(edited.read_bytes(), b"researcher polish")

    def test_low_resolution_image_is_shown_smaller_not_stretched(self):
        Image.new("RGB", (300, 150)).save(self.task / "figures/small.png")
        block = {"t": "Para", "c": [{"t": "Image", "c": [["", [], []], [{"t": "Str", "c": "Small"}], ["figures/small.png", ""]]}]}
        profile = dict(resolve_fonts(self.tokens), text_top_inset_pt=4)
        _, position, _, dpi = build.prepare_image(block, self.task, self.task, [0, 0, 8, 5], self.tokens, profile)
        self.assertAlmostEqual(dpi, 150)
        self.assertAlmostEqual(position[0] + position[2] / 2, 4)  # Still centered in its slot.

    def test_svg_external_resources_are_rejected(self):
        (self.task / "figures/unsafe.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg"><image href="file:///private.png"/></svg>')
        block = {"t": "Para", "c": [{"t": "Image", "c": [["", [], []], [{"t": "Str", "c": "Unsafe"}], ["figures/unsafe.svg", ""]]}]}
        profile = dict(resolve_fonts(self.tokens), text_top_inset_pt=4)
        with self.assertRaisesRegex(StateError, "external resources"):
            build.prepare_image(block, self.task, self.task, [0, 0, 8, 5], self.tokens, profile)

    def test_publish_rolls_back_on_save_failure(self):
        workspace = self.task / ".build"; workspace.mkdir()
        destination = workspace / "output"; destination.mkdir()
        (destination / "build-report.json").write_text("{}")
        (destination / "presentation.pptx").write_bytes(b"previous")
        staging = workspace / "temporary"; staging.mkdir()
        replace = Path.replace
        def fail(source, target):
            if source == staging:
                raise OSError("simulated disk failure")
            return replace(source, target)
        with patch.object(Path, "replace", fail), self.assertRaises(OSError):
            build.publish(staging, self.task, {})
        self.assertEqual((destination / "presentation.pptx").read_bytes(), b"previous")

    def test_publish_does_not_replace_unmanaged_user_directory(self):
        workspace = self.task / ".build"; workspace.mkdir()
        destination = workspace / "output"; destination.mkdir()
        (destination / "user.txt").write_text("user file")
        staging = workspace / "temporary"; staging.mkdir()
        with self.assertRaisesRegex(StateError, "not a managed"):
            build.publish(staging, self.task, {})
        self.assertEqual((destination / "user.txt").read_text(), "user file")

    def test_render_failure_preserves_previous_build(self):
        workspace = self.task / ".build/output"; workspace.mkdir(parents=True)
        (workspace / "build-report.json").write_text("{}")
        (workspace / "presentation.pptx").write_bytes(b"previous")
        with patch.object(build, "assemble", return_value={}), patch.object(build, "render_and_check", side_effect=StateError("render failure")):
            with self.assertRaisesRegex(StateError, "render failure"):
                build.build(self.task)
        self.assertEqual((workspace / "presentation.pptx").read_bytes(), b"previous")
        self.assertFalse((self.task / ".build/.lock").exists())

    def test_input_change_during_render_preserves_previous_build(self):
        workspace = self.task / ".build/output"; workspace.mkdir(parents=True)
        (workspace / "build-report.json").write_text("{}")
        (workspace / "presentation.pptx").write_bytes(b"previous")
        def changed(*args):
            (self.task / "sources/synthetic.csv").write_text("changed")
            return []
        with patch.object(build, "assemble", return_value={}), patch.object(build, "render_and_check", side_effect=changed):
            with self.assertRaisesRegex(StateError, "Inputs or design changed"):
                build.build(self.task)
        self.assertEqual((workspace / "presentation.pptx").read_bytes(), b"previous")

    def test_preview_rejects_stale_inputs(self):
        directory = self.task / ".build/output"; directory.mkdir(parents=True)
        report = {"fingerprint": build.fingerprint(self.task, self.tokens), "artifacts": {}, "outputs": {}, "detected_overflows": 0}
        (directory / "build-report.json").write_text(json.dumps(report))
        self.edit(4, lambda s: s + "\nChanged input\n")
        with self.assertRaisesRegex(StateError, "stale"):
            build.preview(self.task)

    def test_invalid_build_report_has_clear_error(self):
        directory = self.task / ".build/output"; directory.mkdir(parents=True)
        (directory / "build-report.json").write_text("{}")
        with self.assertRaisesRegex(StateError, "Invalid build report"):
            build.preview(self.task)

    def test_unmanaged_previous_directory_is_preserved(self):
        directory = self.task / ".build/.previous"; directory.mkdir(parents=True)
        (directory / "user.txt").write_text("user data")
        staging = self.task / ".build/staging"; staging.mkdir()
        with self.assertRaisesRegex(StateError, "not a managed"):
            build.publish(staging, self.task, {})
        self.assertEqual((directory / "user.txt").read_text(), "user data")

    def test_math_preview_raster_matches_its_box_width(self):
        profile = dict(resolve_fonts(self.tokens), text_top_inset_pt=4)
        ast = json.loads(markdown.pandoc(["-f", "markdown", "-t", "json"], "- 청크 길이 $k$를 8로 둔다"))
        output = self.task / "preview.png"
        height = build.compose(ast["blocks"][0], 9.09, 20, self.task, profile, output=output)
        with Image.open(output) as image:
            self.assertEqual(image.size, (round(9.09 * 300), round(height * 300)))

    def test_fixed_text_measurement_rejects_oversized_tokens(self):
        profile = dict(resolve_fonts(self.tokens), text_top_inset_pt=4)
        block = {"t": "Para", "c": [{"t": "Str", "c": "W" * 200}]}
        with self.assertRaisesRegex(StateError, "fixed content width"):
            build.compose(block, 2, 24, self.task, profile)

    def test_content_height_overflow_is_rejected_before_publication(self):
        api, slides, tokens = markdown.parse_task(self.task)
        slide = self.slide(slides, "S01")
        slide["slots"] = [[{"t": "Para", "c": [{"t": "Str", "c": "Synthetic paragraph for overflow verification. " * 80}]}]]
        profile = dict(resolve_fonts(tokens), text_top_inset_pt=4)
        output = self.task / ".build"; output.mkdir()
        with self.assertRaisesRegex(StateError, "Content overflow"):
            build.assemble(api, [slide], tokens, profile, self.task, output)
        self.assertFalse((output / "presentation.pptx").exists())

    def test_render_check_detects_missing_clipped_text(self):
        import pymupdf
        output = self.task / ".build"; output.mkdir()
        def fake_render(source, target):
            with pymupdf.open() as document:
                page = document.new_page(width=720, height=540)
                page.insert_text((50, 200), "Visible elsewhere", fontsize=24)
                document.save(target)
        report = {"slide_count": 1, "records": [{"page": 1, "slide_id": "S", "role": "title", "box": [.69, .10, 3.43, .56], "expected": "Clipped title", "max_lines": 1}]}
        with patch.object(build, "render_pdf", side_effect=fake_render):
            findings = build.render_and_check(output, report)
        self.assertTrue(any(f["issue"] == "Text missing, altered or clipped" for f in findings))

    def test_native_math_tables_notes_and_branding_round_trip(self):
        api, slides, tokens = markdown.parse_task(self.task)
        slides = slides[:7]  # Cover, contents, section and native content need no browser or LibreOffice.
        profile = dict(resolve_fonts(tokens), text_top_inset_pt=4)
        output = self.task / ".build"; output.mkdir()
        source_hash = sha256(ROOT / tokens["source"]["path"])
        report = build.assemble(api, slides, tokens, profile, self.task, output)
        self.assertEqual(report["native_tables"], 1)
        self.assertEqual(report["native_equations"], 4)
        self.assertEqual(report["notes_slides"], 4)
        presentation = Presentation(output / "presentation.pptx")
        cover = {p.placeholder_format.idx: p.text_frame.text for p in presentation.slides[0].placeholders}
        self.assertEqual(cover, {0: "합성 자료로 제작 기능을 검증한다", 1: "RVLAB 제작·QA 검증", 2: "홍길동 (Gildong Hong)", 3: "경기대학교 창의공과대학 전자공학과 N학년\nRobotics and Vision Lab Robot Team", 4: "발표일시 : 2026.10.11"})
        self.assertIn("Chapter 1: 기능 검증", presentation.slides[1].placeholders[1].text_frame.text)
        top_bar = presentation.slides[3].shapes.title
        self.assertEqual(top_bar.text_frame.text, "표와 수식")
        self.assertEqual([round(v / 914400, 2) for v in (top_bar.left, top_bar.top, top_bar.width, top_bar.height)], tokens["geometry_in"]["title"])
        table = next(s.table for s in presentation.slides[3].shapes if s.has_table)
        table.cell(1, 1).text = "99"
        presentation.slides[3].notes_slide.notes_text_frame.text = "Edited synthetic note"
        target = output / "edited.pptx"; presentation.save(target)
        reopened = Presentation(target)
        self.assertEqual(next(s.table for s in reopened.slides[3].shapes if s.has_table).cell(1, 1).text, "99")
        self.assertEqual(reopened.slides[3].notes_slide.notes_text_frame.text, "Edited synthetic note")
        self.assertEqual(sum(len(list(s._element.iter(f"{{{build.MATH_NS}}}oMath"))) for s in reopened.slides), 4)
        self.assertEqual(sha256(ROOT / tokens["source"]["path"]), source_hash)
        build.package_checks(target, tokens)

    def test_table_preserves_author_marked_best_value(self):
        api, slides, tokens = markdown.parse_task(self.task)
        slide = self.slide(slides, "S01")
        ast = json.loads(markdown.pandoc(["-f", "markdown", "-t", "json"], "| Condition | Synthetic value |\n|---|---|\n| A | **24** |\n"))
        slide["slots"] = [ast["blocks"]]
        profile = dict(resolve_fonts(tokens), text_top_inset_pt=4)
        output = self.task / ".build"; output.mkdir()
        build.assemble(api, [slide], tokens, profile, self.task, output)
        presentation = Presentation(output / "presentation.pptx")
        table = next(s.table for s in presentation.slides[0].shapes if s.has_table)
        self.assertTrue(table.cell(1, 1).text_frame.paragraphs[0].runs[0].font.bold)

    def test_four_figure_slots_stay_equal_when_source_is_reserved(self):
        api, slides, tokens = markdown.parse_task(self.task)
        slide = self.slide(slides, "S08")
        slide["source"] = [{"t": "Para", "c": [{"t": "Str", "c": "Synthetic source"}]}]
        profile = dict(resolve_fonts(tokens), text_top_inset_pt=4)
        output = self.task / ".build"; output.mkdir()
        report = build.assemble(api, [slide], tokens, profile, self.task, output)
        captions = [r["box"] for r in report["records"] if r["role"] == "caption"]
        headline = next(r["box"] for r in report["records"] if r["role"] == "headline")
        first_row_height = captions[0][1] + captions[0][3] - (headline[1] + headline[3] + tokens["geometry_in"]["gap"] / 2)
        second_row_height = captions[2][1] + captions[2][3] - (captions[0][1] + captions[0][3] + .24)
        self.assertAlmostEqual(first_row_height, second_row_height)

    def test_unsupported_native_math_is_reported_as_final_image_fallback(self):
        api, slides, tokens = markdown.parse_task(self.task)
        slide = self.slide(slides, "S02")
        slide["slots"] = [[{"t": "Para", "c": [{"t": "Math", "c": [{"t": "DisplayMath"}, r"\unicode{65}"]}]}]]
        profile = dict(resolve_fonts(tokens), text_top_inset_pt=4)
        output = self.task / ".build"; output.mkdir()
        report = build.assemble(api, [slide], tokens, profile, self.task, output)
        self.assertEqual(report["native_equations"], 0)
        self.assertEqual(report["math_fallbacks"][0]["scope"], "final_and_preview")

    def test_ordered_lists_keep_numbering_and_nested_font_size(self):
        api, slides, tokens = markdown.parse_task(self.task)
        slide = self.slide(slides, "S03"); slide["layout"] = "content"
        ast = json.loads(markdown.pandoc(["-f", "markdown", "-t", "json"], "1. First\n   - Nested\n2. Second\n"))
        slide["slots"] = [ast["blocks"]]
        profile = dict(resolve_fonts(tokens), text_top_inset_pt=4)
        output = self.task / ".build"; output.mkdir()
        build.assemble(api, [slide], tokens, profile, self.task, output)
        from rvppt.template import NS
        presentation = Presentation(output / "presentation.pptx")
        shape = next(s for s in presentation.slides[0].shapes if s.has_text_frame and "First" in s.text)
        self.assertEqual(len(shape._element.findall("p:txBody/a:p/a:pPr/a:buAutoNum", NS)), 2)
        self.assertEqual(shape._element.findall("p:txBody/a:p", NS)[1].find("a:r/a:rPr", NS).get("sz"), "1800")


if __name__ == "__main__":
    unittest.main()
