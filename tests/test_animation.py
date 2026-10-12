import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation

from rvppt import build, markdown
from rvppt.preview import resolve_fonts
from rvppt.state import StateError
from rvppt.template import load_tokens


class AnimationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.task = Path(self.temporary.name)
        (self.task / "figures").mkdir()
        self.source = self.task / "figures/motion.gif"
        self.poster = self.source.with_suffix(".png")
        self.frames = [Image.new("RGB", (600, 300), color) for color in ("white", "gray", "black")]
        self.frames[0].save(self.source, save_all=True, append_images=self.frames[1:], duration=[100, 200, 300], loop=0)
        self.frames[-1].save(self.poster)
        self.block = {"t": "Para", "c": [{"t": "Image", "c": [["", [], []],
                      [{"t": "Str", "c": "합성 애니메이션"}], ["figures/motion.gif", ""]]}]}
        self.tokens = load_tokens()

    def tearDown(self):
        self.temporary.cleanup()

    def test_animation_metadata_and_lint(self):
        metadata = markdown.animation_metadata(self.task, self.source)
        self.assertEqual(metadata, {"path": "figures/motion.gif", "poster": "figures/motion.png",
                                   "frame_count": 3, "duration_ms": 600, "size_px": [600, 300], "loop_count": 0})
        markdown.validate_blocks([self.block], self.task, self.tokens)

    def test_missing_poster_is_rejected_by_lint(self):
        self.poster.unlink()
        with self.assertRaisesRegex(StateError, "matching PNG poster"):
            markdown.validate_blocks([self.block], self.task, self.tokens)

    def test_poster_dimensions_must_match(self):
        Image.new("RGB", (600, 200)).save(self.poster)
        with self.assertRaisesRegex(StateError, "identical pixel dimensions"):
            markdown.animation_metadata(self.task, self.source)

    def test_static_gif_is_rejected(self):
        self.frames[0].save(self.source)
        with self.assertRaisesRegex(StateError, "at least two"):
            markdown.animation_metadata(self.task, self.source)

    def test_frames_need_positive_duration(self):
        self.frames[0].save(self.source, save_all=True, append_images=self.frames[1:], duration=[100, 0, 300])
        with self.assertRaisesRegex(StateError, "frame 2 requires a positive duration"):
            markdown.animation_metadata(self.task, self.source)

    def test_corrupt_gif_is_reported_as_input_error(self):
        self.source.write_bytes(b"GIF89a")
        with self.assertRaisesRegex(StateError, "Invalid GIF or PNG"):
            markdown.animation_metadata(self.task, self.source)

    def test_poster_symlink_is_rejected(self):
        self.poster.unlink()
        self.poster.symlink_to(self.source)
        with self.assertRaisesRegex(StateError, "symlinks"):
            markdown.animation_metadata(self.task, self.source)

    def test_oversized_gif_is_rejected_before_decoding(self):
        with self.source.open("r+b") as file:
            file.truncate(64 * 1024 * 1024 + 1)
        with self.assertRaisesRegex(StateError, "64 MiB"):
            markdown.animation_metadata(self.task, self.source)

    def test_math_compatible_deck_preserves_original_gif(self):
        ast = json.loads(markdown.pandoc(["-f", "markdown", "-t", "json"], "$$x^2$$"))
        output = self.task / "build"
        output.mkdir()
        spec = {"id": "S01", "layout": "content", "topic": "애니메이션", "headline": "수식과 상태 변화를 확인한다",
                "slots": [[ast["blocks"][0], self.block]], "source": [], "notes": []}
        profile = dict(resolve_fonts(self.tokens), text_top_inset_pt=4)
        build.assemble(ast["pandoc-api-version"], [spec], self.tokens, profile, self.task, output)
        with ZipFile(output / "presentation-math-compatible.pptx") as package:
            media = [name for name in package.namelist() if name.endswith(".gif")]
            self.assertEqual(len(media), 1)
            self.assertEqual(package.read(media[0]), self.source.read_bytes())
        compatible = Presentation(output / "presentation-math-compatible.pptx")
        self.assertEqual(len(list(compatible.slides[0]._element.iter(f"{{{build.MATH_NS}}}oMath"))), 0)

    def test_final_gif_bytes_and_selected_static_preview_are_independent(self):
        output = self.task / "build"; output.mkdir()
        original_gif, original_poster = self.source.read_bytes(), self.poster.read_bytes()
        spec = {"id": "S01", "layout": "figure_2", "topic": "애니메이션", "headline": "변화의 순서를 확인한다",
                "slots": [[self.block], [self.block]], "source": [], "notes": []}
        profile = dict(resolve_fonts(self.tokens), text_top_inset_pt=4)
        report = build.assemble([], [spec], self.tokens, profile, self.task, output)
        with ZipFile(output / "presentation.pptx") as package:
            media = [name for name in package.namelist() if name.endswith(".gif")]
            self.assertEqual(len(media), 1)
            self.assertEqual(package.read(media[0]), original_gif)
        with ZipFile(output / "presentation-preview-copy.pptx") as package:
            self.assertFalse(any(name.endswith(".gif") for name in package.namelist()))
        final = Presentation(output / "presentation.pptx")
        preview = Presentation(output / "presentation-preview-copy.pptx")
        animated = [shape for shape in final.slides[0].shapes if shape.shape_type == 13]
        static = [shape for shape in preview.slides[0].shapes if shape.shape_type == 13]
        self.assertEqual(len(animated), 2)
        self.assertEqual(len(static), 2)
        for motion, still in zip(animated, static):
            self.assertEqual(motion.image.blob, original_gif)
            self.assertEqual(still.image.blob, original_poster)
            self.assertEqual((motion.left, motion.top, motion.width, motion.height),
                             (still.left, still.top, still.width, still.height))
        self.assertEqual(report["images"][0]["animation"]["duration_ms"], 600)
        self.assertGreaterEqual(report["images"][0]["effective_dpi"], self.tokens["limits"]["image_min_effective_dpi"])
        self.assertEqual(self.source.read_bytes(), original_gif)
        self.assertEqual(self.poster.read_bytes(), original_poster)


if __name__ == "__main__":
    unittest.main()
